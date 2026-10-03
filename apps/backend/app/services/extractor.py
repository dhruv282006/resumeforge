"""Document text extraction pipeline for PDF and DOCX files."""

from __future__ import annotations

import io
from abc import ABC, abstractmethod


class DocumentExtractor(ABC):
    """Abstract document text extractor."""

    @abstractmethod
    def extract(self, file_bytes: bytes) -> str:
        """Extract plain text from file bytes."""
        ...


class PdfExtractor(DocumentExtractor):
    """PDF text extractor using pypdf."""

    def extract(self, file_bytes: bytes) -> str:
        """Extract text from all pages in PDF bytes."""
        import pypdf

        stream = io.BytesIO(file_bytes)
        try:
            reader = pypdf.PdfReader(stream)
            extracted_pages: list[str] = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                cleaned = text.strip()
                if cleaned:
                    extracted_pages.append(cleaned)
            return "\n\n".join(extracted_pages)
        except Exception as exc:
            raise ValueError(f"Failed to extract text from PDF: {exc}") from exc


class DocxExtractor(DocumentExtractor):
    """DOCX text extractor using python-docx."""

    def extract(self, file_bytes: bytes) -> str:
        """Extract paragraphs and table text from DOCX bytes."""
        import docx

        stream = io.BytesIO(file_bytes)
        try:
            doc = docx.Document(stream)
            extracted_elements: list[str] = []

            for paragraph in doc.paragraphs:
                cleaned = paragraph.text.strip()
                if cleaned:
                    extracted_elements.append(cleaned)

            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        extracted_elements.append(" | ".join(row_texts))

            return "\n\n".join(extracted_elements)
        except Exception as exc:
            raise ValueError(f"Failed to extract text from DOCX: {exc}") from exc


def get_extractor(filename_or_mime: str) -> DocumentExtractor:
    """Factory to retrieve appropriate document extractor by extension or mime."""
    lower = filename_or_mime.lower()
    if lower.endswith(".pdf") or "application/pdf" in lower:
        return PdfExtractor()
    elif (
        lower.endswith(".docx")
        or "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in lower
    ):
        return DocxExtractor()
    else:
        raise ValueError(
            f"Unsupported document format: '{filename_or_mime}'. Only PDF and DOCX files are supported."
        )

