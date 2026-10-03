"""Tests for Profile and Resume REST API endpoints."""

from __future__ import annotations

import io
import pytest
from docx import Document
from httpx import AsyncClient


class TestProfileApiEndpoints:
    @pytest.mark.asyncio
    async def test_get_and_put_profile(self, client_with_db: AsyncClient) -> None:
        # GET /api/v1/profile
        res = await client_with_db.get("/api/v1/profile")
        assert res.status_code == 200
        data = res.json()
        assert data["full_name"] == "Candidate User"

        # PUT /api/v1/profile
        update_payload = {
            "headline": "Lead Systems Architect",
            "location": "Seattle, WA",
            "years_of_experience": 8.5,
            "preferred_roles": ["Staff Engineer", "Principal Architect"],
        }
        res_put = await client_with_db.put("/api/v1/profile", json=update_payload)
        assert res_put.status_code == 200
        updated = res_put.json()
        assert updated["headline"] == "Lead Systems Architect"
        assert updated["location"] == "Seattle, WA"
        assert "Principal Architect" in updated["preferred_roles"]

    @pytest.mark.asyncio
    async def test_profile_completeness_endpoint(self, client_with_db: AsyncClient) -> None:
        res = await client_with_db.get("/api/v1/profile/completeness")
        assert res.status_code == 200
        data = res.json()
        assert "overall" in data
        assert "sections" in data
        assert isinstance(data["overall"], int)


class TestResumeApiEndpoints:
    @pytest.mark.asyncio
    async def test_resume_crud_flow(self, client_with_db: AsyncClient) -> None:
        # 1. Create resume
        create_res = await client_with_db.post(
            "/api/v1/resumes",
            json={
                "title": "Full Stack Engineer Resume",
                "target_role": "Full Stack Developer",
                "initial_version_number": "v1.0",
            },
        )
        assert create_res.status_code == 201
        resume = create_res.json()
        resume_id = resume["id"]
        assert resume["title"] == "Full Stack Engineer Resume"
        assert len(resume["versions"]) == 1

        # 2. Get resume
        get_res = await client_with_db.get(f"/api/v1/resumes/{resume_id}")
        assert get_res.status_code == 200
        assert get_res.json()["id"] == resume_id

        # 3. Clone version
        v1_id = resume["versions"][0]["id"]
        clone_res = await client_with_db.post(
            f"/api/v1/resumes/{resume_id}/versions/{v1_id}/clone",
            json={"version_number": "v1.1", "change_summary": "Added React 19 section"},
        )
        assert clone_res.status_code == 201
        assert clone_res.json()["version_number"] == "v1.1"

        # 4. Activate new version
        v1_1_id = clone_res.json()["id"]
        act_res = await client_with_db.post(
            f"/api/v1/resumes/{resume_id}/versions/{v1_1_id}/activate"
        )
        assert act_res.status_code == 200
        assert act_res.json()["active_version_id"] == v1_1_id

    @pytest.mark.asyncio
    async def test_upload_resume_file_endpoint(self, client_with_db: AsyncClient) -> None:
        # Create resume
        res = await client_with_db.post(
            "/api/v1/resumes",
            json={"title": "Data Scientist Resume", "initial_version_number": "v1.0"},
        )
        resume_id = res.json()["id"]

        # Create docx in memory
        doc = Document()
        doc.add_heading("Senior Data Scientist", 0)
        doc.add_paragraph("Machine learning modeling with PyTorch and Transformers.")
        buf = io.BytesIO()
        doc.save(buf)
        file_bytes = buf.getvalue()

        # Upload DOCX
        upload_res = await client_with_db.post(
            f"/api/v1/resumes/{resume_id}/versions/upload",
            files={"file": ("ds_resume.docx", file_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"version_number": "v2.0", "change_summary": "DOCX resume upload"},
        )
        assert upload_res.status_code == 201
        uploaded_version = upload_res.json()
        assert uploaded_version["version_number"] == "v2.0"
        assert uploaded_version["status"] == "READY"
        assert "Senior Data Scientist" in uploaded_version["extracted_text"]
        assert "PyTorch and Transformers" in uploaded_version["extracted_text"]

