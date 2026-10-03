"""Tests for ResumeService and ResumeVersion lifecycle operations."""

from __future__ import annotations

import io
import uuid
import pytest
from docx import Document
from sqlalchemy.ext.asyncio import AsyncSession

from app.schemas.resume import (
    ResumeCreate,
    ResumeUpdate,
    ResumeVersionClone,
    ResumeVersionCreate,
)
from app.services.resume_service import ResumeService
from app.services.storage import LocalStorageProvider


class TestResumeService:
    @pytest.mark.asyncio
    async def test_create_and_list_resumes(self, db_session: AsyncSession, tmp_path) -> None:
        storage = LocalStorageProvider(base_dir=tmp_path)
        service = ResumeService(db_session, storage=storage)
        user_id = uuid.uuid4()

        resume = await service.create_resume(
            ResumeCreate(
                title="Backend Developer Resume",
                target_role="Senior Backend Engineer",
                target_company="Stripe",
                initial_version_number="v1.0",
            ),
            user_id=user_id,
        )
        assert resume.title == "Backend Developer Resume"
        assert resume.target_role == "Senior Backend Engineer"
        assert len(resume.versions) == 1
        assert resume.versions[0].version_number == "v1.0"
        assert resume.active_version_id == resume.versions[0].id

        all_resumes = await service.list_resumes(user_id)
        assert len(all_resumes) == 1
        assert all_resumes[0].id == resume.id

    @pytest.mark.asyncio
    async def test_version_cloning_and_immutability(self, db_session: AsyncSession, tmp_path) -> None:
        storage = LocalStorageProvider(base_dir=tmp_path)
        service = ResumeService(db_session, storage=storage)
        user_id = uuid.uuid4()

        resume = await service.create_resume(
            ResumeCreate(title="ML Engineer Resume", initial_version_number="v1.0"),
            user_id=user_id,
        )
        v1 = resume.versions[0]

        # Clone v1.0 into v1.1
        cloned = await service.clone_version(
            resume.id,
            v1.id,
            ResumeVersionClone(
                version_number="v1.1",
                change_summary="Tailored for distributed systems role",
            ),
            user_id=user_id,
        )
        assert cloned is not None
        assert cloned.version_number == "v1.1"
        assert cloned.parent_version_id == v1.id

        versions = await service.list_versions(resume.id, user_id)
        assert len(versions) == 2
        # Original version remains intact
        assert any(v.version_number == "v1.0" for v in versions)
        assert any(v.version_number == "v1.1" for v in versions)

    @pytest.mark.asyncio
    async def test_activate_version(self, db_session: AsyncSession, tmp_path) -> None:
        storage = LocalStorageProvider(base_dir=tmp_path)
        service = ResumeService(db_session, storage=storage)
        user_id = uuid.uuid4()

        resume = await service.create_resume(
            ResumeCreate(title="Cloud Architect Resume", initial_version_number="v1.0"),
            user_id=user_id,
        )
        v1 = resume.versions[0]

        v2 = await service.create_version(
            resume.id,
            ResumeVersionCreate(version_number="v2.0", change_summary="Major revamp"),
            user_id=user_id,
        )
        assert v2 is not None

        # Activate v2.0
        updated_resume = await service.activate_version(resume.id, v2.id, user_id)
        assert updated_resume is not None
        assert updated_resume.active_version_id == v2.id

    @pytest.mark.asyncio
    async def test_upload_and_extract_docx_version(self, db_session: AsyncSession, tmp_path) -> None:
        storage = LocalStorageProvider(base_dir=tmp_path)
        service = ResumeService(db_session, storage=storage)
        user_id = uuid.uuid4()

        resume = await service.create_resume(
            ResumeCreate(title="Software Engineer Resume", initial_version_number="v1.0"),
            user_id=user_id,
        )

        # Build real docx in memory
        doc = Document()
        doc.add_heading("DevOps & Cloud Engineer", 0)
        doc.add_paragraph("Expert in Kubernetes, Docker, Terraform, and AWS.")
        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        version = await service.upload_and_process_version(
            resume_id=resume.id,
            file_bytes=docx_bytes,
            filename="devops_resume.docx",
            user_id=user_id,
            version_number="v1.1",
            change_summary="Uploaded primary DOCX resume",
        )
        assert version.status == "READY"
        assert version.version_number == "v1.1"
        assert "DevOps & Cloud Engineer" in version.extracted_text
        assert "Kubernetes, Docker, Terraform" in version.extracted_text
        assert version.file_path is not None
        assert await storage.exists(version.file_path) is True

